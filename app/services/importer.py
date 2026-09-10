import json
import re
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models import Company, Contact, ContactRoute
from app.services.email_quality import validate_contact_email
from app.services.no_go_companies import company_match_key, is_no_go_company
from app.time_utils import utc_now

EMAIL_RE = re.compile(r"[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}", re.I)
LINKEDIN_RE = re.compile(r"https?://[^\s,;]*linkedin\.com/[^\s,;]+", re.I)
URL_RE = re.compile(r"https?://[^\s,;]+", re.I)
PHONE_RE = re.compile(r"(?:\+?\d[\d\s()./-]{6,}\d)")
REJECTED_OUTREACH_LOCAL_RE = re.compile(r"(^|[._+-])(careers?|hire|hr|jobs?|recruit(?:ment|ing)?)([._+-]|$)", re.I)

GENERIC_EMAIL_LOCALS = {
    "admin",
    "contact",
    "customerservice",
    "enquiries",
    "enquiry",
    "hello",
    "info",
    "mail",
    "office",
    "parts",
    "sales",
    "service",
    "support",
    "team",
    "webteam",
}

HEADER_ALIASES = {
    "company": ["公司名", "公司", "公司名称", "标准公司名", "company", "company name", "name"],
    "country": ["国家", "国家地区", "国家/地区", "标准国家", "country"],
    "region": ["区域", "地区", "省州", "省/州", "region", "subregion", "sub-region", "territory", "area"],
    "website": ["公司官网", "官网", "网站", "website", "company website", "url"],
    "priority": ["评级", "当前评级", "客户等级", "priority", "rating"],
    "company_type": ["公司类型", "类型", "company_type", "company type"],
    "intro": ["公司介绍", "介绍", "company introduction", "company description", "description", "profile", "about"],
    "source": ["来源", "标准来源", "source"],
    "notes": ["备注", "标准备注", "跟进备注", "notes"],
    "full_name": ["联系人", "姓名", "原始联系人", "contact", "contact name", "full_name", "full name", "name"],
    "position": ["联系人职位", "职位", "原始联系人职位", "position", "title"],
    "route": ["联系方式", "联系信息", "联系方法", "原始联系方式", "contact route", "contact routes"],
    "email": ["邮箱", "电子邮箱", "email", "e-mail", "mail"],
    "phone": ["电话", "手机号", "手机", "telephone", "phone", "mobile"],
    "linkedin": ["linkedin", "linkedin url", "linkedin profile"],
    "whatsapp": ["whatsapp", "whats app"],
    "follow_record": ["跟进记录", "跟进情况", "标准跟进记录", "follow record", "follow-up record"],
    "follow_idea": ["跟进思路", "follow idea", "follow-up idea"],
}

CORE_COMPANY_COLUMNS = {"company"}
ROUTE_COLUMNS = {"route", "email", "phone", "linkedin", "whatsapp"}
MIN_HEADER_HITS = 3
FILL_DOWN_COLUMNS = {"company", "country", "region", "website", "priority", "company_type", "intro", "source", "notes"}


@dataclass
class SheetPlan:
    title: str
    header_row: int
    columns: dict[str, int]


def _clean_cell(value: object) -> str:
    text = str(value or "").strip()
    if text.lower() in {"nan", "none", "null"} or text in {"\\", "/", "-", "暂无"}:
        return ""
    return text


def _source_label(filename: str) -> str:
    return Path(filename or "uploaded_bd_excel").stem or "uploaded_bd_excel"


def _is_generic_email(email: str) -> bool:
    local = email.split("@", 1)[0].lower()
    if local in GENERIC_EMAIL_LOCALS:
        return True
    return any(local.startswith(f"{prefix}.") or local.startswith(f"{prefix}-") or local.startswith(f"{prefix}_") for prefix in GENERIC_EMAIL_LOCALS)


def is_generic_email(email: str) -> bool:
    return _is_generic_email(email)


def _normalize_email(value: str) -> str | None:
    check = validate_contact_email(value, check_deliverability=False)
    if check.is_blocking:
        return None
    return (check.normalized_email or value.strip().lower()).strip(".,;:").lower()


def _is_rejected_outreach_email(email: str) -> bool:
    local = email.split("@", 1)[0].lower()
    return bool(REJECTED_OUTREACH_LOCAL_RE.search(local))


def _as_int(value: object) -> int | None:
    text = str(value or "").strip()
    if not text:
        return None
    try:
        return int(float(text))
    except ValueError:
        return None


def _merge_source(existing: str | None, new_value: str) -> str:
    values: list[str] = []
    for chunk in (existing or "").replace("\n", ";").split(";"):
        cleaned = chunk.strip()
        if cleaned and cleaned not in values:
            values.append(cleaned)
    for chunk in new_value.replace("\n", ";").split(";"):
        cleaned = chunk.strip()
        if cleaned and cleaned not in values:
            values.append(cleaned)
    return "; ".join(values)


def _text_len(value: str | None) -> int:
    return len((value or "").strip())


def _header_key(value: object) -> str:
    text = str(value or "").strip().lower()
    text = re.sub(r"\s+", " ", text)
    return text


def _find_sheet_plan(ws) -> SheetPlan | None:
    alias_lookup = {
        alias.lower(): canonical
        for canonical, aliases in HEADER_ALIASES.items()
        for alias in aliases
    }
    best: tuple[int, dict[str, int], int] | None = None
    for row_idx in range(1, min(ws.max_row, 30) + 1):
        columns: dict[str, int] = {}
        for col_idx in range(1, ws.max_column + 1):
            header = _header_key(ws.cell(row_idx, col_idx).value)
            canonical = alias_lookup.get(header)
            if canonical and canonical not in columns:
                columns[canonical] = col_idx
        score = len(columns)
        if best is None or score > best[2]:
            best = (row_idx, columns, score)
    if not best:
        return None
    row_idx, columns, score = best
    if score < MIN_HEADER_HITS or not CORE_COMPANY_COLUMNS.issubset(columns) or not (set(columns) & ROUTE_COLUMNS):
        return None
    return SheetPlan(ws.title, row_idx, columns)


def _merged_value_lookup(ws) -> dict[tuple[int, int], object]:
    lookup: dict[tuple[int, int], object] = {}
    for merged_range in ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = merged_range.bounds
        value = ws.cell(min_row, min_col).value
        for row_idx in range(min_row, max_row + 1):
            for col_idx in range(min_col, max_col + 1):
                lookup[(row_idx, col_idx)] = value
    return lookup


def _cell_text(ws, merged_lookup: dict[tuple[int, int], object], row_idx: int, col_idx: int | None) -> str:
    if not col_idx:
        return ""
    value = ws.cell(row_idx, col_idx).value
    if value is None:
        value = merged_lookup.get((row_idx, col_idx))
    return _clean_cell(value)


def _extract_routes(raw_route: str) -> dict[str, list[str]]:
    emails: list[str] = []
    for email in EMAIL_RE.findall(raw_route):
        normalized = _normalize_email(email)
        if normalized and normalized not in emails:
            emails.append(normalized)

    linkedins: list[str] = []
    for url in LINKEDIN_RE.findall(raw_route):
        cleaned = url.strip().rstrip(".,;")
        if cleaned not in linkedins:
            linkedins.append(cleaned)

    phones: list[str] = []
    without_links = EMAIL_RE.sub(" ", URL_RE.sub(" ", raw_route))
    for phone in PHONE_RE.findall(without_links):
        cleaned = re.sub(r"\s+", " ", phone).strip(" /,;")
        digits = re.sub(r"\D", "", cleaned)
        if len(digits) >= 7 and cleaned not in phones:
            phones.append(cleaned)

    has_whatsapp_hint = bool(re.search(r"\bwhatsapp\b|wa\.me", raw_route, re.I))
    return {
        "email": emails,
        "linkedin": linkedins,
        "whatsapp": phones if has_whatsapp_hint else [],
        "phone": [] if has_whatsapp_hint else phones,
    }


def _combined_route_text(row_data: dict[str, str]) -> str:
    values = [
        row_data.get("route", ""),
        row_data.get("email", ""),
        row_data.get("phone", ""),
        row_data.get("linkedin", ""),
        row_data.get("whatsapp", ""),
    ]
    return "\n".join(value for value in values if value)


def _company_notes(row_data: dict[str, str], source: str, sheet_name: str, row_idx: int) -> str:
    parts = []
    labels = [
        ("公司官网", row_data.get("website", "")),
        ("公司介绍", row_data.get("intro", "")),
        ("备注", row_data.get("notes", "")),
        ("跟进记录", row_data.get("follow_record", "")),
        ("跟进思路", row_data.get("follow_idea", "")),
    ]
    for label, value in labels:
        if value:
            parts.append(f"{label}: {value}")
    parts.append(f"导入来源: {source}")
    parts.append(f"原始位置: {sheet_name}!{row_idx}")
    return "\n".join(parts)


def _json_metadata_notes(row_data: dict[str, str], batch_id: str, source_system: str, selection_rule: str) -> str:
    labels = [
        ("bd_company_id", row_data.get("bd_company_id", "")),
        ("domain", row_data.get("domain", "")),
        ("batch_id", batch_id),
        ("source_system", source_system),
        ("selection_rule", selection_rule),
        ("priority_dynamic_score", row_data.get("priority_dynamic_score", "")),
        ("priority_roll_status", row_data.get("priority_roll_status", "")),
        ("priority_contact_status", row_data.get("priority_contact_status", "")),
        ("priority_contact_level", row_data.get("priority_contact_level", "")),
    ]
    return "\n".join(f"{key}: {value}" for key, value in labels if value)


def _get_or_create_company(session: Session, row_data: dict[str, str], source: str, sheet_name: str, row_idx: int) -> tuple[Company, bool, bool]:
    company_name = row_data.get("company", "").strip()
    company = session.exec(select(Company).where(Company.name == company_name)).first()
    notes = _company_notes(row_data, source, sheet_name, row_idx)
    if company is None:
        company = Company(
            name=company_name,
            country=row_data.get("country") or None,
            region=row_data.get("region") or None,
            company_type=row_data.get("company_type") or None,
            source=source,
            priority=row_data.get("priority") or "B",
            status="new",
            notes=notes or None,
        )
        session.add(company)
        session.flush()
        return company, True, False

    changed = False
    if row_data.get("country") and not company.country:
        company.country = row_data["country"]
        changed = True
    if row_data.get("region") and not company.region:
        company.region = row_data["region"]
        changed = True
    if row_data.get("company_type") and not company.company_type:
        company.company_type = row_data["company_type"]
        changed = True
    if row_data.get("priority") and (not company.priority or company.priority == "B"):
        company.priority = row_data["priority"]
        changed = True
    merged_source = _merge_source(company.source, source)
    if merged_source != (company.source or ""):
        company.source = merged_source
        changed = True
    existing_notes = company.notes or ""
    if notes and notes not in existing_notes:
        if "bd_company_id:" in notes or "batch_id:" in notes:
            company.notes = f"{existing_notes}\n\n{notes}".strip()
        elif _text_len(notes) > _text_len(existing_notes):
            company.notes = notes
        changed = True
    if changed:
        company.updated_at = utc_now()
        session.add(company)
    return company, False, changed


def _find_contact_by_identity(session: Session, company_id: int, full_name: str, position: str) -> Contact | None:
    contacts = session.exec(select(Contact).where(Contact.company_id == company_id, Contact.full_name == full_name)).all()
    if not contacts:
        return None
    if position:
        for contact in contacts:
            if (contact.position or "").strip().lower() == position.strip().lower():
                return contact
    return contacts[0]


def _get_or_create_contact(
    session: Session,
    company: Company,
    full_name: str,
    position: str,
    primary_email: str,
    priority_contact_rank: int | None = None,
) -> tuple[Contact, bool, bool]:
    contact = session.exec(select(Contact).where(Contact.email == primary_email)).first()
    if contact is None and company.id is not None and full_name and full_name != "Team":
        contact = _find_contact_by_identity(session, company.id, full_name, position)
    if contact is None:
        has_contact = session.exec(select(Contact).where(Contact.company_id == company.id)).first() is not None
        contact = Contact(
            company_id=company.id,
            full_name=full_name,
            first_name=full_name.split(" ")[0] if full_name != "Team" else "Team",
            position=position or None,
            email=primary_email,
            priority_contact_rank=priority_contact_rank,
            is_primary=not has_contact,
            status="active",
        )
        session.add(contact)
        session.flush()
        return contact, True, False

    changed = False
    if contact.company_id != company.id:
        raise ValueError(f"Email {primary_email} already belongs to another company; resolve the conflict before importing.")
    if full_name and contact.full_name != full_name and contact.full_name == "Team":
        contact.full_name = full_name
        contact.first_name = full_name.split(" ")[0]
        changed = True
    if position and not contact.position:
        contact.position = position
        changed = True
    if priority_contact_rank is not None and contact.priority_contact_rank != priority_contact_rank:
        contact.priority_contact_rank = priority_contact_rank
        changed = True
    if changed:
        contact.updated_at = utc_now()
        session.add(contact)
    return contact, False, changed


def _upsert_route(
    session: Session,
    contact: Contact,
    route_type: str,
    route_value: str,
    source: str,
    raw_cell: str,
    is_primary: bool = False,
) -> tuple[bool, bool]:
    route = session.exec(
        select(ContactRoute).where(
            ContactRoute.contact_id == contact.id,
            ContactRoute.route_type == route_type,
            ContactRoute.route_value == route_value,
        )
    ).first()
    if route is None:
        route = ContactRoute(
            contact_id=contact.id,
            route_type=route_type,
            route_value=route_value,
            source=source,
            raw_cell=raw_cell or None,
            is_primary=is_primary,
            status="active",
        )
        session.add(route)
        return True, False

    changed = False
    merged_source = _merge_source(route.source, source)
    if merged_source != (route.source or ""):
        route.source = merged_source
        changed = True
    if raw_cell and not route.raw_cell:
        route.raw_cell = raw_cell
        changed = True
    if is_primary and not route.is_primary:
        route.is_primary = True
        changed = True
    if changed:
        route.updated_at = utc_now()
        session.add(route)
    return False, changed


def _new_report(source_label: str, valid_sheets: int = 0, skipped_sheets: int = 0) -> dict[str, int | list[int] | str | dict[str, int]]:
    return {
        "created_companies": 0,
        "created_contacts": 0,
        "updated_companies": 0,
        "updated_contacts": 0,
        "created_routes": 0,
        "updated_routes": 0,
        "duplicates": 0,
        "invalid_email": 0,
        "skipped_no_email": 0,
        "generic_contacts": 0,
        "processed_rows": 0,
        "valid_sheets": valid_sheets,
        "skipped_sheets": skipped_sheets,
        "source_label": source_label,
        "contact_ids": [],
        "skipped": {},
    }


def _skip(report: dict[str, int | list[int] | str | dict[str, int]], reason: str) -> None:
    skipped = report["skipped"]
    if isinstance(skipped, dict):
        skipped[reason] = skipped.get(reason, 0) + 1


def _is_blacklisted_company(company: Company | None, row_data: dict[str, str]) -> bool:
    if company is not None and is_no_go_company(company):
        return True
    return str(row_data.get("blacklist") or row_data.get("is_blacklisted") or "").strip().lower() in {"1", "true", "yes", "blacklisted"}


def _import_normalized_rows(
    session: Session,
    rows: list[dict[str, str]],
    source_label: str,
    valid_sheets: int = 0,
    skipped_sheets: int = 0,
    raise_on_empty: bool = True,
) -> dict[str, int | list[int] | str | dict[str, int]]:
    report = _new_report(source_label, valid_sheets=valid_sheets, skipped_sheets=skipped_sheets)
    contact_ids: list[int] = report["contact_ids"]  # type: ignore[assignment]
    no_go_companies = {
        company_match_key(company.name): company
        for company in session.exec(select(Company)).all()
        if is_no_go_company(company)
    }

    for row_data in rows:
        skip_reason = row_data.get("_skip_reason", "")
        if skip_reason:
            _skip(report, skip_reason)
            continue

        company_name = row_data.get("company", "").strip()
        raw_route = _combined_route_text(row_data).strip()
        if not company_name and not raw_route:
            continue
        if not company_name:
            _skip(report, "missing_company")
            continue

        existing_company = session.exec(select(Company).where(Company.name == company_name)).first()
        if company_match_key(company_name) in no_go_companies or _is_blacklisted_company(existing_company, row_data):
            _skip(report, "blacklisted_company")
            continue

        routes = _extract_routes(raw_route)
        source = row_data.get("source") or source_label
        blocked_emails = [email for email in routes["email"] if _is_rejected_outreach_email(email)]
        emails = [email for email in routes["email"] if email not in blocked_emails]
        if blocked_emails:
            _skip(report, "hire_email")

        sheet_name = row_data.get("_sheet_name", "normalized")
        row_idx = _as_int(row_data.get("_row_idx")) or 0

        if not emails:
            if routes["email"] or row_data.get("email"):
                report["invalid_email"] = int(report["invalid_email"]) + 1
                _skip(report, "invalid_email")
            if existing_company is not None and existing_company.id is not None:
                full_name = row_data.get("full_name", "")
                existing_contact = (
                    _find_contact_by_identity(session, existing_company.id, full_name, row_data.get("position", ""))
                    if full_name
                    else None
                )
                if existing_contact is not None:
                    for linkedin in routes["linkedin"]:
                        route_created, route_updated = _upsert_route(session, existing_contact, "linkedin", linkedin, source, raw_route)
                        report["created_routes"] = int(report["created_routes"]) + int(route_created)
                        report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)
                    for whatsapp in routes["whatsapp"]:
                        route_created, route_updated = _upsert_route(session, existing_contact, "whatsapp", whatsapp, source, raw_route)
                        report["created_routes"] = int(report["created_routes"]) + int(route_created)
                        report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)
                    for phone in routes["phone"]:
                        route_created, route_updated = _upsert_route(session, existing_contact, "phone", phone, source, raw_route)
                        report["created_routes"] = int(report["created_routes"]) + int(route_created)
                        report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)
            report["skipped_no_email"] = int(report["skipped_no_email"]) + 1
            continue

        report["processed_rows"] = int(report["processed_rows"]) + 1
        company, company_created, company_updated = _get_or_create_company(session, row_data, source, sheet_name, row_idx)
        report["created_companies"] = int(report["created_companies"]) + int(company_created)
        report["updated_companies"] = int(report["updated_companies"]) + int(company_updated)

        generic_emails = [email for email in emails if _is_generic_email(email)]
        personal_emails = [email for email in emails if email not in generic_emails]
        contact_groups: list[tuple[str, list[str]]] = []
        if personal_emails:
            contact_groups.append(("person", personal_emails))
        if generic_emails:
            contact_groups.append(("team", generic_emails))

        priority_contact_rank = _as_int(row_data.get("priority_contact_rank"))
        ready_for_prep = row_data.get("_ready_for_prep", "1") != "0"
        for group_type, group_emails in contact_groups:
            if group_type == "team":
                full_name = "Team"
                position = "General inbox"
                group_priority_contact_rank = None
                report["generic_contacts"] = int(report["generic_contacts"]) + 1
            else:
                full_name = row_data.get("full_name") or "Team"
                position = row_data.get("position", "")
                group_priority_contact_rank = priority_contact_rank
            primary_email = group_emails[0]
            contact, contact_created, contact_updated = _get_or_create_contact(
                session,
                company,
                full_name,
                position,
                primary_email,
                priority_contact_rank=group_priority_contact_rank,
            )
            report["created_contacts"] = int(report["created_contacts"]) + int(contact_created)
            report["updated_contacts"] = int(report["updated_contacts"]) + int(contact_updated)
            if ready_for_prep and contact.id is not None and contact.id not in contact_ids:
                contact_ids.append(contact.id)

            for idx, email in enumerate(group_emails):
                route_created, route_updated = _upsert_route(session, contact, "email", email, source, raw_route, is_primary=(email == contact.email or idx == 0))
                report["created_routes"] = int(report["created_routes"]) + int(route_created)
                report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)
                report["duplicates"] = int(report["duplicates"]) + int(not route_created and not route_updated)
            for linkedin in routes["linkedin"]:
                route_created, route_updated = _upsert_route(session, contact, "linkedin", linkedin, source, raw_route)
                report["created_routes"] = int(report["created_routes"]) + int(route_created)
                report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)
            for whatsapp in routes["whatsapp"]:
                route_created, route_updated = _upsert_route(session, contact, "whatsapp", whatsapp, source, raw_route)
                report["created_routes"] = int(report["created_routes"]) + int(route_created)
                report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)
            for phone in routes["phone"]:
                route_created, route_updated = _upsert_route(session, contact, "phone", phone, source, raw_route)
                report["created_routes"] = int(report["created_routes"]) + int(route_created)
                report["updated_routes"] = int(report["updated_routes"]) + int(route_updated)

    if int(report["processed_rows"]) == 0 and raise_on_empty:
        raise ValueError("鏈湪鑱旂郴鏂瑰紡鍒椾腑鎵惧埌鍙鍏ラ偖绠憋紝鏈啓鍏ュ鎴峰簱銆?")

    session.commit()
    return report


def _row_values(
    ws,
    merged_lookup: dict[tuple[int, int], object],
    plan: SheetPlan,
    row_idx: int,
    carry: dict[str, str],
) -> dict[str, str]:
    row_data = {
        canonical: _cell_text(ws, merged_lookup, row_idx, col_idx)
        for canonical, col_idx in plan.columns.items()
    }
    company_name = row_data.get("company", "").strip()
    if company_name and company_name.casefold() != carry.get("company", "").strip().casefold():
        carry.clear()
    for key in FILL_DOWN_COLUMNS:
        if row_data.get(key):
            carry[key] = row_data[key]
        elif carry.get(key):
            row_data[key] = carry[key]
    return row_data


def import_contacts(session: Session, payload: bytes, filename: str) -> dict[str, int | list[int] | str | dict[str, int]]:
    suffix = Path(filename.lower()).suffix
    if suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("Please upload an Excel file (.xlsx or .xlsm).")

    try:
        workbook = load_workbook(BytesIO(payload), data_only=True)
    except Exception as exc:
        raise ValueError(f"Unable to read Excel file: {exc}") from exc

    plans = [(ws, _find_sheet_plan(ws)) for ws in workbook.worksheets]
    valid = [(ws, plan) for ws, plan in plans if plan is not None]
    if not valid:
        raise ValueError("No importable sheet found. A company column and at least one contact-route column are required.")

    source_fallback = _source_label(filename)
    normalized_rows: list[dict[str, str]] = []
    for ws, plan in valid:
        merged_lookup = _merged_value_lookup(ws)
        carry: dict[str, str] = {}
        for row_idx in range(plan.header_row + 1, ws.max_row + 1):
            row_data = _row_values(ws, merged_lookup, plan, row_idx, carry)
            row_data["_sheet_name"] = plan.title
            row_data["_row_idx"] = str(row_idx)
            normalized_rows.append(row_data)

    return _import_normalized_rows(
        session,
        normalized_rows,
        source_fallback,
        valid_sheets=len(valid),
        skipped_sheets=len(plans) - len(valid),
        raise_on_empty=True,
    )


def _json_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _json_company_key(row: dict[str, str]) -> str:
    return (
        row.get("bd_company_id")
        or row.get("domain")
        or row.get("company")
        or ""
    ).strip().lower()


def _validate_json_rank_uniqueness(rows: list[dict[str, str]]) -> None:
    seen: dict[tuple[str, int], int] = {}
    duplicates: list[str] = []
    for idx, row in enumerate(rows, start=1):
        rank = _as_int(row.get("priority_contact_rank"))
        if rank is None:
            continue
        key = _json_company_key(row)
        if not key:
            continue
        marker = (key, rank)
        if marker in seen:
            duplicates.append(f"row {seen[marker]} and row {idx}: {key} rank {rank}")
        else:
            seen[marker] = idx
    if duplicates:
        raise ValueError("Duplicate priority_contact_rank in the same company scope: " + "; ".join(duplicates[:10]))


def _normalize_bd_json_rows(data: dict) -> tuple[str, list[dict[str, str]]]:
    batch_id = _json_text(data.get("batch_id"))
    source_system = _json_text(data.get("source_system"))
    selection_rule = _json_text(data.get("selection_rule"))
    contacts = data.get("contacts")
    if not batch_id or not source_system or not isinstance(contacts, list):
        raise ValueError("JSON must include batch_id, source_system and contacts[].")

    source_label = f"{source_system}:{batch_id}"
    rows: list[dict[str, str]] = []
    for idx, item in enumerate(contacts, start=1):
        if not isinstance(item, dict):
            rows.append({"_skip_reason": "invalid_contact_object", "_row_idx": str(idx), "_sheet_name": "bd-database-json"})
            continue

        row = {str(key): _json_text(value) for key, value in item.items()}
        row["_sheet_name"] = "bd-database-json"
        row["_row_idx"] = str(idx)
        row["source"] = row.get("source") or source_label
        row["region"] = row.get("region") or row.get("company_region") or row.get("区域") or row.get("地区") or ""
        row["notes"] = "\n".join(
            part
            for part in [
                row.get("notes", ""),
                _json_metadata_notes(row, batch_id, source_system, selection_rule),
            ]
            if part
        )

        email_send_status = row.get("email_send_status", "").strip().lower()
        if email_send_status and email_send_status != "include":
            row["_skip_reason"] = "email_send_status"
        elif not row.get("company") or not row.get("email") or not row.get("full_name"):
            row["_skip_reason"] = "missing_required"
        elif row.get("priority_roll_status", "").strip().lower() != "ready":
            row["_skip_reason"] = "not_ready"
        elif _is_rejected_outreach_email(row.get("email", "")):
            row["_skip_reason"] = "hire_email"
        else:
            row["_ready_for_prep"] = "1"

        row["route"] = "\n".join(part for part in [row.get("email", ""), row.get("linkedin", ""), row.get("whatsapp", "")] if part)
        rows.append(row)

    _validate_json_rank_uniqueness([row for row in rows if not row.get("_skip_reason")])
    return source_label, rows


def import_bd_json_candidates(session: Session, payload: bytes, filename: str = "bd-database.json") -> dict[str, int | list[int] | str | dict[str, int]]:
    suffix = Path(filename.lower()).suffix
    if suffix != ".json":
        raise ValueError("Please upload a .json file.")
    try:
        data = json.loads(payload.decode("utf-8-sig"))
    except Exception as exc:
        raise ValueError(f"Unable to read JSON file: {exc}") from exc
    if not isinstance(data, dict):
        raise ValueError("JSON root must be an object.")

    source_label, rows = _normalize_bd_json_rows(data)
    report = _import_normalized_rows(
        session,
        rows,
        source_label,
        valid_sheets=1,
        skipped_sheets=0,
        raise_on_empty=False,
    )
    report["batch_id"] = _json_text(data.get("batch_id"))
    report["source_system"] = _json_text(data.get("source_system"))
    report["selection_rule"] = _json_text(data.get("selection_rule"))
    return report

