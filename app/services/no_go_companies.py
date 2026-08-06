import csv
import re
from io import BytesIO, StringIO
from pathlib import Path

from openpyxl import load_workbook
from sqlmodel import Session, select

from app.models import Company
from app.time_utils import utc_now


NO_GO_STATUSES = {"blacklist", "blacklisted", "blocked", "suppressed"}
COMPANY_HEADER_ALIASES = {
    "company",
    "company name",
    "company_name",
    "name",
    "公司",
    "公司名称",
    "企业",
    "企业名称",
}


def normalize_company_name(value: object) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def company_match_key(value: object) -> str:
    return normalize_company_name(value).casefold()


def is_no_go_company(company: Company) -> bool:
    return (company.status or "").strip().lower() in NO_GO_STATUSES


def _decode_text(payload: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030", "cp1252"):
        try:
            return payload.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("无法识别文件编码，请另存为 UTF-8 CSV 或 TXT 后重试。")


def _names_from_rows(rows: list[list[object]]) -> list[str]:
    populated_rows = [
        row
        for row in rows
        if any(normalize_company_name(value) for value in row)
    ]
    if not populated_rows:
        return []

    first_row = populated_rows[0]
    company_column = 0
    data_start = 0
    for index, value in enumerate(first_row):
        if company_match_key(value) in COMPANY_HEADER_ALIASES:
            company_column = index
            data_start = 1
            break

    names: list[str] = []
    for row in populated_rows[data_start:]:
        if company_column >= len(row):
            continue
        name = normalize_company_name(row[company_column])
        if name:
            names.append(name)
    return names


def parse_no_go_company_text(value: str) -> list[str]:
    rows = [[line] for line in value.splitlines()]
    return _names_from_rows(rows)


def parse_no_go_company_file(payload: bytes, filename: str) -> list[str]:
    suffix = Path(filename).suffix.lower()
    if suffix == ".txt":
        return parse_no_go_company_text(_decode_text(payload))
    if suffix == ".csv":
        text = _decode_text(payload)
        try:
            dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [list(row) for row in csv.reader(StringIO(text), dialect)]
        return _names_from_rows(rows)
    if suffix in {".xlsx", ".xlsm"}:
        workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
        names: list[str] = []
        try:
            for worksheet in workbook.worksheets:
                rows = [list(row) for row in worksheet.iter_rows(values_only=True)]
                names.extend(_names_from_rows(rows))
        finally:
            workbook.close()
        return names
    raise ValueError("仅支持 .xlsx、.xlsm、.csv 或 .txt NO-GO 公司名单。")


def import_no_go_companies(
    session: Session,
    names: list[str],
    source_label: str,
) -> dict[str, int]:
    companies = session.exec(select(Company)).all()
    company_by_key: dict[str, Company] = {}
    for company in companies:
        company_by_key.setdefault(company_match_key(company.name), company)

    result = {
        "input": 0,
        "created": 0,
        "updated": 0,
        "already_no_go": 0,
        "duplicates": 0,
    }
    seen: set[str] = set()
    for raw_name in names:
        name = normalize_company_name(raw_name)
        key = company_match_key(name)
        if not key:
            continue
        result["input"] += 1
        if key in seen:
            result["duplicates"] += 1
            continue
        seen.add(key)

        company = company_by_key.get(key)
        if company is None:
            company = Company(
                name=name,
                priority="C",
                status="blacklisted",
                source=source_label,
                notes=f"NO-GO company imported from {source_label}.",
            )
            session.add(company)
            company_by_key[key] = company
            result["created"] += 1
            continue
        if is_no_go_company(company):
            result["already_no_go"] += 1
            continue

        company.status = "blacklisted"
        company.updated_at = utc_now()
        session.add(company)
        result["updated"] += 1

    if not seen:
        raise ValueError("名单中没有找到公司名称。请每行填写一家，或使用带 company/company name/公司名称 列的文件。")
    session.commit()
    return result
