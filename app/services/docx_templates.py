import base64
import re
from html import escape
from io import BytesIO

from docx import Document
from docx.text.run import Run
from docx.oxml.ns import qn

EMU_PER_PIXEL = 9525
TAG_RE = re.compile(r"<[^>]+>")
JINJA_VAR_RE = re.compile(r"\{\{.*?\}\}", re.S)


def _image_size_style(blip, *, max_width: str = "100%") -> str:
    drawing = blip
    while drawing is not None and drawing.tag != qn("w:drawing"):
        drawing = drawing.getparent()
    extent = drawing.xpath(".//wp:extent")[0] if drawing is not None and drawing.xpath(".//wp:extent") else None
    if extent is None:
        return f"max-width:{max_width};height:auto"
    try:
        width = max(int(int(extent.get("cx")) / EMU_PER_PIXEL), 1)
        height = max(int(int(extent.get("cy")) / EMU_PER_PIXEL), 1)
    except (TypeError, ValueError):
        return f"max-width:{max_width};height:auto"
    return f"width:{width}px;height:{height}px;max-width:{max_width}"


def _run_image_html(run, *, display: str = "block", margin: str = "12px 0") -> str:
    html_parts = []
    for blip in run._element.xpath(".//a:blip"):
        relationship_id = blip.get(qn("r:embed"))
        if not relationship_id:
            continue
        relationship = run.part.rels.get(relationship_id)
        if relationship is None:
            continue
        image_part = relationship.target_part
        if not image_part.content_type.startswith("image/"):
            continue
        payload_b64 = base64.b64encode(image_part.blob).decode("ascii")
        size_style = _image_size_style(blip)
        html_parts.append(
            '<img src="data:{content_type};base64,{payload}" '
            'style="{size_style};display:{display};margin:{margin};vertical-align:middle;" alt="" />'.format(
                content_type=image_part.content_type,
                payload=payload_b64,
                size_style=size_style,
                display=display,
                margin=margin,
            )
        )
    return "".join(html_parts)


def _run_style(run) -> str:
    styles = []
    if run.font.size:
        styles.append(f"font-size:{run.font.size.pt:g}pt")
    if run.font.color and run.font.color.rgb:
        styles.append(f"color:#{run.font.color.rgb}")
    return ";".join(styles)


def _style_run(text: str, *, bold: bool, italic: bool, underline: bool) -> str:
    value = escape(text).replace("\xa0", "&nbsp;").replace("\n", "<br>")
    if not value:
        return ""
    if underline:
        value = f"<u>{value}</u>"
    if italic:
        value = f"<em>{value}</em>"
    if bold:
        value = f"<strong>{value}</strong>"
    return value


def _run_to_html(run, *, image_display: str = "block", image_margin: str = "12px 0") -> str:
    text_html = _style_run(
        run.text,
        bold=bool(run.bold),
        italic=bool(run.italic),
        underline=bool(run.underline),
    )
    style = _run_style(run)
    if text_html and style:
        text_html = f'<span style="{style}">{text_html}</span>'
    return f"{text_html}{_run_image_html(run, display=image_display, margin=image_margin)}"


def _paragraph_child_runs(paragraph):
    for child in paragraph._p.iterchildren():
        if child.tag == qn("w:r"):
            yield Run(child, paragraph), None
        elif child.tag == qn("w:hyperlink"):
            relationship_id = child.get(qn("r:id"))
            href = ""
            if relationship_id and relationship_id in paragraph.part.rels:
                href = paragraph.part.rels[relationship_id].target_ref
            for run_element in child.iterchildren(qn("w:r")):
                yield Run(run_element, paragraph), href or None


def _paragraph_runs_to_html(paragraph, *, image_display: str = "block", image_margin: str = "12px 0") -> str:
    html_parts = []
    for run, href in _paragraph_child_runs(paragraph):
        run_html = _run_to_html(run, image_display=image_display, image_margin=image_margin)
        if href and run_html:
            run_html = f'<a href="{escape(href, quote=True)}">{run_html}</a>'
        html_parts.append(run_html)
    return "".join(html_parts)


def _paragraph_style(paragraph) -> str:
    styles = []
    if paragraph.alignment is not None:
        alignment_map = {
            0: "left",
            1: "center",
            2: "right",
            3: "justify",
        }
        alignment = alignment_map.get(int(paragraph.alignment))
        if alignment:
            styles.append(f"text-align:{alignment}")
    return ";".join(styles)


def _paragraph_to_html(paragraph) -> str:
    html = _paragraph_runs_to_html(paragraph)
    style = _paragraph_style(paragraph)
    style_attr = f' style="{style}"' if style else ""
    if not html.strip():
        return f"<p{style_attr}><br></p>"
    return f"<p{style_attr}>{html}</p>"


def _repair_split_jinja_placeholders(html: str) -> str:
    def clean(match: re.Match[str]) -> str:
        value = TAG_RE.sub("", match.group(0))
        value = re.sub(r"\s+", " ", value)
        return value.strip()

    return JINJA_VAR_RE.sub(clean, html)


def _paragraph_to_text(paragraph) -> str:
    return paragraph.text.replace("\xa0", " ").strip()


def _paragraph_has_image(paragraph) -> bool:
    return any(run._element.xpath(".//a:blip") for run in paragraph.runs)


def _paragraph_has_content(paragraph) -> bool:
    return bool(paragraph.text.strip()) or _paragraph_has_image(paragraph)


def _paragraph_image_count(paragraph) -> int:
    return sum(len(run._element.xpath(".//a:blip")) for run in paragraph.runs)


def _paragraph_to_signature_html(paragraph) -> str:
    image_count = _paragraph_image_count(paragraph)
    text = paragraph.text.replace("\xa0", "").strip()
    if image_count >= 2 and not text:
        html = _paragraph_runs_to_html(paragraph, image_display="inline-block", image_margin="0 6px 0 0")
        return f'<div class="bd-signature-social" style="margin:10px 0 14px 0;white-space:nowrap;">{html}</div>'
    if image_count == 1 and not text:
        html = _paragraph_runs_to_html(paragraph, image_display="block", image_margin="0")
        return f'<div class="bd-signature-logo" style="margin:8px 0 0 0;">{html}</div>'
    html = _paragraph_runs_to_html(paragraph, image_display="inline-block", image_margin="0 6px 0 0")
    style = _paragraph_style(paragraph)
    style_attr = f' style="{style}"' if style else ""
    if not html.strip():
        return f"<p{style_attr}><br></p>"
    return f"<p{style_attr}>{html}</p>"


def parse_docx_email_template(payload: bytes, fallback_subject: str | None = None) -> tuple[str, str, str]:
    document = Document(BytesIO(payload))
    paragraphs = [paragraph for paragraph in document.paragraphs if _paragraph_has_content(paragraph)]
    if not paragraphs:
        raise ValueError("DOCX template is empty.")

    first_line = _paragraph_to_text(paragraphs[0])
    if first_line.lower().startswith("subject:"):
        subject = first_line.split(":", 1)[1].strip()
        body_paragraphs = paragraphs[1:]
    else:
        subject = (fallback_subject or "").strip() or first_line
        body_paragraphs = paragraphs[1:] if not (fallback_subject or "").strip() else paragraphs
    if not subject:
        raise ValueError("Subject line is empty.")
    if not body_paragraphs:
        raise ValueError("DOCX template body is empty.")

    body_html = _repair_split_jinja_placeholders("\n".join(_paragraph_to_html(paragraph) for paragraph in body_paragraphs))
    body_text = "\n\n".join(_paragraph_to_text(paragraph) for paragraph in body_paragraphs)
    return subject, body_html, body_text


def parse_docx_signature_html(payload: bytes, include_text: bool = False) -> str:
    document = Document(BytesIO(payload))
    html_parts = []
    if include_text:
        paragraphs = [paragraph for paragraph in document.paragraphs if _paragraph_has_content(paragraph)]
        html_parts.extend(_paragraph_to_signature_html(paragraph) for paragraph in paragraphs)
    else:
        image_paragraphs = [paragraph for paragraph in document.paragraphs if _paragraph_has_image(paragraph)]
        html_parts.extend(_paragraph_to_signature_html(paragraph) for paragraph in image_paragraphs)
    if not html_parts:
        raise ValueError("DOCX signature template is empty.")
    return "\n".join(html_parts)
